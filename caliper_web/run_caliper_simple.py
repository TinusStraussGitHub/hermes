#!/usr/bin/env python3
"""Non-GUI Caliper processor — uses user's exact logic, applies settings."""
import os, re, sys
from openpyxl import load_workbook

folder = sys.argv[1]
template_path = sys.argv[2]
output_path = sys.argv[3]

client = os.environ.get('CALIPER_CLIENT','BME Client')
mine = os.environ.get('CALIPER_MINE','Delmas')
block_id = os.environ.get('CALIPER_BLOCK','Block-A')
planned_d = float(os.environ.get('CALIPER_PLANNED_D','165'))
final_stem = float(os.environ.get('CALIPER_FINAL_STEM','1'))
ave_density = float(os.environ.get('CALIPER_DENSITY','1.2'))
operator = os.environ.get('CALIPER_OPERATOR','')
date_cal = os.environ.get('CALIPER_DATE','')

wb = load_workbook(template_path)
info = wb['Information']
info['C1']=client; info['C2']=mine; info['C3']=block_id
info['C4']=planned_d; info['C5']=final_stem; info['C6']=ave_density
info['C7']=operator; info['C8']=date_cal

las_files = [f for f in os.listdir(folder) if f.lower().endswith('.las')]
used_sheets = []

for i,file_name in enumerate(las_files):
    file_path = os.path.join(folder, file_name)
    with open(file_path,'r') as f:
        lines = f.readlines()
    lines = lines[20:]  # skip header
    data = []
    for line in lines:
        cols = re.split(r'\s+', line.strip())
        if cols and cols[0]:
            data.append(cols)
    sheet_name = f'Sheet{i+1}'
    if sheet_name not in wb.sheetnames:
        continue
    sheet = wb[sheet_name]
    used_sheets.append(sheet_name)
    for row_num, cols in enumerate(data, start=2):
        for col_num, val in enumerate(cols, start=1):
            sheet.cell(row=row_num, column=col_num).value = val
    # Name in L1
    sheet.cell(row=1, column=12).value = os.path.splitext(file_name)[0]
    # Convert depth/diameter
    for row_num in range(2, 2+len(data)):
        for col in [1,2]:
            v = sheet.cell(row=row_num, column=col).value
            try:
                sheet.cell(row=row_num, column=col).value = float(v)
            except:
                pass
    # Delete negative diameter rows + bottom 50
    last_row = sheet.max_row
    start_row = max(2, last_row - 49)
    rows_to_delete = last_row - start_row + 1
    row = 2
    while row <= sheet.max_row:
        d_val = sheet.cell(row=row, column=2).value
        if isinstance(d_val, float) and d_val < 0:
            sheet.delete_rows(row, amount=rows_to_delete)
        else:
            row += 1
    # Delete depth < final stemming
    row = 2
    while row <= sheet.max_row:
        depth_val = sheet.cell(row=row, column=1).value
        if isinstance(depth_val, float) and depth_val < final_stem:
            sheet.delete_rows(row)
        else:
            row += 1
    # Format A-H to 2 decimals
    for row_num in range(2, sheet.max_row+1):
        for col in range(1,9):
            cell = sheet.cell(row=row_num, column=col)
            cell.number_format = '0.00'
    # Sort by depth (column A)
    max_row = sheet.max_row
    if max_row > 2:
        sorted_data = []
        for r in range(2, max_row+1):
            a = sheet.cell(row=r, column=1).value
            b = sheet.cell(row=r, column=2).value
            if isinstance(a, float):
                sorted_data.append((a,b))
        sorted_data.sort(key=lambda x: x[0])
        for idx, (a,b) in enumerate(sorted_data, start=2):
            sheet.cell(row=idx, column=1).value = a
            sheet.cell(row=idx, column=2).value = b
    # Insert formulas C-H
    max_row = sheet.max_row
    last_data_row = 0
    for r in range(2, max_row+1):
        if sheet.cell(row=r, column=1).value is not None:
            last_data_row = r
    if last_data_row > 1:
        formulas = {
            'C': '=B{row}/2',
            'D': '=(B{row}-(B{row}*2))/2',
            'E': '=G{row}/2',
            'F': '=(G{row}-(G{row}*2))/2',
            'G': '=Information!C$4',
            'H': '=B{row}-G{row}'
        }
        for r in range(2, last_data_row+1):
            for col_letter, formula in formulas.items():
                sheet[f'{col_letter}{r}'].value = formula.format(row=r)

# Delete unused sheets (preserve core)
core_sheets = ['Graphs','Table','Information']
deleted_sheets = [s for s in wb.sheetnames if s not in used_sheets and s not in core_sheets]
for s in deleted_sheets:
    del wb[s]

# Update Table
if 'Table' in wb.sheetnames:
    table = wb['Table']
    imported = []
    for s in used_sheets:
        val = wb[s].cell(row=1, column=12).value
        if val:
            imported.append(val)
    for i, name in enumerate(imported):
        if i < 50:
            table.cell(row=i+3, column=1).value = name
    for row in range(52, 2, -1):
        val = table.cell(row=row, column=1).value
        if val not in imported:
            table.delete_rows(row)
    # Find min/ave/max rows
    min_r = ave_r = max_r = None
    for r in range(3, table.max_row+1):
        v = table.cell(row=r, column=1).value
        if v == 'Minimum': min_r = r
        elif v == 'Average': ave_r = r
        elif v == 'Maximum': max_r = r
    cols_min = ['B','C','D','E','G','H','K','L','M','N','O']
    cols_ave = ['B','C','D','E','G','I','K','L','M','N','O']
    cols_max = ['B','C','D','E','G','J','K','L','M','N','O']
    if min_r:
        for c in cols_min:
            table[f"{c}{min_r}"].value = f"=MIN({c}3:{c}{min_r-1})"
    if ave_r:
        for c in cols_ave:
            table[f"{c}{ave_r}"].value = f"=AVERAGE({c}3:{c}{ave_r-1})"
    if max_r:
        for c in cols_max:
            table[f"{c}{max_r}"].value = f"=MAX({c}3:{c}{max_r-1})"

wb.save(output_path)
print(f"Processed {len(las_files)} .las files -> {output_path}")
